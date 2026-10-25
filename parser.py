import os
import re
from datetime import datetime

# Common technical and professional skills dictionary for taxonomy matching
SKILLS_DB = {
    # Programming Languages
    "python", "java", "c++", "c#", "c", "javascript", "typescript", "ruby", "php", "go",
    "golang", "rust", "swift", "kotlin", "scala", "r", "dart", "perl", "bash", "shell",
    "sql", "html", "html5", "css", "css3", "sass", "matlab",
    
    # Frameworks & Libraries
    "react", "react.js", "reactjs", "angular", "angularjs", "vue", "vue.js", "vuejs",
    "next.js", "nextjs", "node", "node.js", "nodejs", "express", "express.js",
    "flask", "django", "fastapi", "spring", "spring boot", "asp.net", ".net",
    "bootstrap", "tailwind", "tailwindcss", "redux", "jquery", "flutter",
    "react native", "electron", "pytorch", "tensorflow", "keras", "scikit-learn",
    "sklearn", "pandas", "numpy", "scipy", "opencv", "matplotlib", "seaborn",
    "spacy", "nltk", "huggingface", "llm", "langchain",
    
    # Databases & Storage
    "mysql", "postgresql", "postgres", "sqlite", "mongodb", "redis",
    "cassandra", "oracle", "mariadb", "dynamodb", "elasticsearch", "neo4j",
    "firebase", "supabase", "snowflake", "bigquery",
    
    # Cloud, DevOps & Tools
    "aws", "amazon web services", "azure", "gcp", "google cloud", "docker",
    "kubernetes", "k8s", "terraform", "ansible", "jenkins", "git", "github",
    "gitlab", "ci/cd", "circleci", "linux", "unix", "nginx", "apache",
    "kafka", "rabbitmq", "rest", "rest api", "graphql", "microservices",
    "celery", "grpc",
    
    # Methodologies & Data
    "agile", "scrum", "jira", "unit testing", "pytest", "selenium", "jest",
    "machine learning", "deep learning", "nlp", "computer vision", "data analysis",
    "data science", "data engineering", "power bi", "tableau", "excel",
    "etl", "system design", "devops",
    
    # Professional & Soft Skills
    "leadership", "communication", "team management", "project management",
    "problem solving", "critical thinking", "collaboration", "client relations"
}

DEGREE_PATTERNS = [
    r"\b(?:b\.?tech|bachelor of technology)\b",
    r"\b(?:b\.?e\.?|bachelor of engineering)\b",
    r"\b(?:b\.?s\.?|b\.?sc\.?|bachelor of science)\b",
    r"\b(?:b\.?c\.?a\.?|bachelor of computer applications)\b",
    r"\b(?:b\.?a\.?|bachelor of arts)\b",
    r"\b(?:b\.?com\.?|bachelor of commerce)\b",
    r"\b(?:bachelor'?s(?:\s+degree)?)\b",
    r"\b(?:m\.?tech|master of technology)\b",
    r"\b(?:m\.?e\.?|master of engineering)\b",
    r"\b(?:m\.?s\.?|m\.?sc\.?|master of science)\b",
    r"\b(?:m\.?c\.?a\.?|master of computer applications)\b",
    r"\b(?:m\.?b\.?a\.?|master of business administration)\b",
    r"\b(?:master'?s(?:\s+degree)?)\b",
    r"\b(?:ph\.?d\.?|doctor of philosophy|doctorate)\b",
    r"\b(?:diploma)\b",
    r"\b(?:associate'?s(?:\s+degree)?)\b"
]

def extract_text_from_pdf(filepath):
    """Safely extracts readable text from a PDF file."""
    text_content = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        if len(reader.pages) == 0:
            raise ValueError("The uploaded PDF has 0 pages.")
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    except Exception as e:
        raise ValueError(f"Could not read PDF document: {str(e)}")
        
    full_text = "\n".join(text_content).strip()
    if not full_text:
        raise ValueError("The PDF document does not contain selectable text (it may be a scanned image or empty).")
    return full_text

def extract_text_from_docx(filepath):
    """Safely extracts readable text from a DOCX file including tables."""
    text_content = []
    try:
        from docx import Document
        doc = Document(filepath)
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_content.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))
    except Exception as e:
        raise ValueError(f"Could not read DOCX document: {str(e)}")
        
    full_text = "\n".join(text_content).strip()
    if not full_text:
        raise ValueError("The DOCX document is empty or contains no readable text paragraphs.")
    return full_text

def extract_resume_text(filepath):
    """Extracts raw text from either PDF or DOCX format."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Only .pdf and .docx files are permitted.")

def extract_email(text):
    """Extracts contact email from resume text if present."""
    pattern = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
    matches = re.findall(pattern, text)
    if matches:
        return matches[0].strip().rstrip(".")
    return ""

def extract_phone(text):
    """Extracts contact phone number if present."""
    pattern = r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{2,4}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}"
    lines = text.split("\n")[:25]
    for line in lines:
        candidates = re.findall(pattern, line)
        for c in candidates:
            digits_only = re.sub(r"\D", "", c)
            if 10 <= len(digits_only) <= 15:
                return c.strip()
    return ""

def extract_name(text):
    """
    Extracts candidate name ONLY if reliably identified from top non-header lines.
    Never invents or fabricates fake candidate information.
    """
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        return ""
        
    blacklist_words = {
        "resume", "curriculum", "vitae", "cv", "profile", "contact",
        "email", "phone", "address", "location", "summary", "experience",
        "education", "skills", "projects", "objective", "page", "portfolio",
        "github", "linkedin", "qualification", "qualifications"
    }
    
    # Check top 6 non-empty lines
    for line in lines[:6]:
        # Header or key-value lines with colon are not names
        if ":" in line:
            continue
            
        # Handle lines with separators like "Alex Morgan | Software Engineer" or "Alex Morgan - Developer"
        parts = re.split(r"[\s]*[|\-–—•/][\s]*", line)
        candidate_part = parts[0].strip()
        cleaned = re.sub(r"[^\w\s]", "", candidate_part).strip()
        words = cleaned.split()
        if 2 <= len(words) <= 4:
            lower_words = [w.lower() for w in words]
            # Must not contain header blacklist words
            if any(w in blacklist_words for w in lower_words):
                continue
            # If every word is a recognized tech skill (e.g. "Python Docker"), reject
            if all(w in SKILLS_DB for w in lower_words):
                continue
            # Must not contain digits, email, or website
            if any("@" in w or any(c.isdigit() for c in w) for w in words):
                continue
            # Must look like proper capitalized alphabetical words (e.g. "Sarah Jenkins")
            if all(len(w) >= 2 and w[0].isupper() and w.isalpha() for w in words):
                return " ".join([w.capitalize() for w in words])
                
    return ""

def extract_skills(text):
    """
    Extracts technical and professional skills matching known skill taxonomy.
    """
    found_skills = set()
    text_lower = text.lower()
    
    for skill in SKILLS_DB:
        if skill in ["c++", "c#", ".net", "asp.net"]:
            escaped = re.escape(skill)
            if re.search(r"(?:^|\s|[,/|])" + escaped + r"(?:$|\s|[,/|;])", text_lower):
                found_skills.add(skill)
        elif len(skill) <= 2:
            pattern = r"\b" + re.escape(skill) + r"\b"
            if re.search(pattern, text_lower):
                found_skills.add(skill)
        else:
            pattern = r"(?:^|\W)" + re.escape(skill) + r"(?:\W|$)"
            if re.search(pattern, text_lower):
                if skill in ["react.js", "reactjs"]:
                    found_skills.add("react")
                elif skill in ["node.js", "nodejs"]:
                    found_skills.add("node.js")
                elif skill in ["vue.js", "vuejs"]:
                    found_skills.add("vue")
                elif skill in ["next.js", "nextjs"]:
                    found_skills.add("next.js")
                elif skill in ["golang"]:
                    found_skills.add("go")
                elif skill in ["postgres"]:
                    found_skills.add("postgresql")
                else:
                    found_skills.add(skill)
                    
    return sorted(list(found_skills))

def extract_experience_years(text):
    """
    Estimates candidate's total years of experience from direct statements or date ranges.
    Returns float number of years, or None if not determinable.
    """
    # 1. Direct explicit statement search
    explicit_exp_patterns = [
        r"(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:experience|exp)",
        r"(?:experience|exp)\s*(?::|-)?\s*(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)",
        r"(?:over|around|more than)\s+(\d+(?:\.\d+)?)\s*(?:years?|yrs?)"
    ]
    
    for pat in explicit_exp_patterns:
        match = re.search(pat, text, re.IGNORECASE)
        if match:
            try:
                years = float(match.group(1))
                if 0 <= years <= 40:
                    return years
            except ValueError:
                pass
                
    # 2. Date ranges scan: e.g. 2018 - 2023, 2020 - Present
    current_year = datetime.now().year
    year_range_pattern = r"\b(199\d|200\d|201\d|202\d)\s*(?:-|–|to)\s*(199\d|200\d|201\d|202\d|present|current)\b"
    ranges = re.findall(year_range_pattern, text, re.IGNORECASE)
    
    intervals = []
    for start_str, end_str in ranges:
        start_yr = int(start_str)
        if end_str.lower() in ["present", "current"]:
            end_yr = current_year
        else:
            end_yr = int(end_str)
            
        if 1990 <= start_yr <= current_year and start_yr <= end_yr <= current_year + 1:
            intervals.append((start_yr, end_yr))
            
    if intervals:
        intervals.sort(key=lambda x: x[0])
        merged = []
        for interval in intervals:
            if not merged:
                merged.append(interval)
            else:
                prev_start, prev_end = merged[-1]
                if interval[0] <= prev_end:
                    merged[-1] = (prev_start, max(prev_end, interval[1]))
                else:
                    merged.append(interval)
        total_spanned = sum(end - start for start, end in merged)
        if 0 < total_spanned <= 40:
            return float(total_spanned)
            
    return None

def extract_education(text):
    """Extracts degree names and education institutions found in resume."""
    degrees_found = []
    text_lower = text.lower()
    
    for pattern in DEGREE_PATTERNS:
        matches = re.finditer(pattern, text_lower)
        for m in matches:
            matched_str = m.group(0).strip().upper()
            if matched_str not in degrees_found:
                degrees_found.append(matched_str)
                
    edu_section_pattern = r"(?:education|academics|qualifications)([\s\S]*?)(?:experience|work|projects|skills|certifications|$)"
    match = re.search(edu_section_pattern, text, re.IGNORECASE)
    if match:
        section_lines = [l.strip() for l in match.group(1).split("\n") if l.strip()]
        for l in section_lines[:4]:
            if len(l) > 5 and not any(l == d for d in degrees_found):
                if not re.search(r"^(education|academics)", l, re.IGNORECASE):
                    degrees_found.append(l)
                    
    return degrees_found[:5]

def extract_projects(text):
    """Extracts highlighted projects or initiatives from project section."""
    projects_found = []
    proj_section_pattern = r"(?:projects|academic projects|key projects)([\s\S]*?)(?:experience|education|skills|certifications|awards|$)"
    match = re.search(proj_section_pattern, text, re.IGNORECASE)
    if match:
        section_lines = [l.strip() for l in match.group(1).split("\n") if l.strip()]
        for line in section_lines:
            if line.startswith(("-", "•", "*", "1.", "2.", "3.", "4.", "5.")):
                clean_bullet = re.sub(r"^[-•*0-9\.\s]+", "", line).strip()
                if len(clean_bullet) > 10:
                    projects_found.append(clean_bullet)
            elif 5 < len(line) < 90 and not line.endswith("."):
                projects_found.append(line)
            if len(projects_found) >= 5:
                break
                
    return projects_found

def parse_resume(filepath, original_filename=""):
    """
    Main parser entry point: extracts raw text and parses all entities without fabrication.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")
        
    file_size = os.path.getsize(filepath)
    ext = os.path.splitext(filepath)[1].lower()
    
    raw_text = extract_resume_text(filepath)
    if not raw_text or not raw_text.strip():
        raise ValueError("Could not extract readable text from the file. It may be empty or an image-only scan.")
        
    email = extract_email(raw_text)
    phone = extract_phone(raw_text)
    name = extract_name(raw_text)  # Empty string if not reliably detected
    skills = extract_skills(raw_text)
    experience_years = extract_experience_years(raw_text)
    education = extract_education(raw_text)
    projects = extract_projects(raw_text)
    
    return {
        "name": name,  # Empty string if not found, never fabricated
        "email": email,
        "phone": phone,
        "experience_years": experience_years,  # Can be None if not found
        "skills": skills,
        "education": education,
        "projects": projects,
        "raw_text": raw_text,
        "filename": os.path.basename(filepath),
        "original_filename": original_filename or os.path.basename(filepath),
        "file_type": ext.replace(".", "").upper(),
        "file_size": file_size
    }
